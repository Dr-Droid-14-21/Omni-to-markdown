from __future__ import annotations

from pathlib import Path
import re

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "docs" / "benchmark_reports"
OUT_DIR.mkdir(parents=True, exist_ok=True)

DOCX_PATH = OUT_DIR / "AI_Model_Benchmark_Comparison_2026.docx"
XLSX_PATH = OUT_DIR / "AI_Model_Benchmark_Comparison_2026.xlsx"
CHART_DIR = OUT_DIR / "chart_images"
CHART_DIR.mkdir(parents=True, exist_ok=True)

MODELS = [
    "ChatGPT 5.5",
    "ChatGPT 5.3 Codex",
    "Claude Opus 4.7",
    "Claude Sonnet 4.6",
    "MiniMax M2.7",
    "Kimi K2.6",
    "DeepSeek V4-Pro Max",
    "GLM-5.1",
]

MODEL_NOTES = {
    "ChatGPT 5.5": "OpenAI's April 2026 frontier model for ChatGPT, Codex, documents, spreadsheets, and long-running agentic work.",
    "ChatGPT 5.3 Codex": "OpenAI's February 2026 Codex model, tuned for coding agents and computer-based professional work.",
    "Claude Opus 4.7": "Anthropic's April 2026 flagship Claude model for deep reasoning, long-horizon coding, and knowledge work.",
    "Claude Sonnet 4.6": "Anthropic's February 2026 lower-cost Claude model, positioned as near-Opus quality for many coding and agent tasks.",
    "MiniMax M2.7": "MiniMax's open model aimed at agentic coding, office productivity, and self-improving workflows.",
    "Kimi K2.6": "Moonshot AI's open-weight multimodal agentic model with strong coding, search, and reasoning coverage.",
    "DeepSeek V4-Pro Max": "DeepSeek's latest V4-Pro model in maximum reasoning mode; this matches the user's 'DeepSeek4Pro latest' description.",
    "GLM-5.1": "Z.ai's open-weight agentic model with broad coding, math, tool-use, and long-horizon benchmark coverage.",
}

SOURCES = [
    {
        "id": "openai_gpt55",
        "name": "OpenAI - Introducing GPT-5.5",
        "url": "https://openai.com/index/introducing-gpt-5-5/",
        "note": "ChatGPT 5.5 / GPT-5.5 coding, agentic, knowledge-work, math, and cyber scores.",
    },
    {
        "id": "openai_gpt53_codex",
        "name": "OpenAI - Introducing GPT-5.3-Codex",
        "url": "https://openai.com/index/introducing-gpt-5-3-codex/",
        "note": "ChatGPT 5.3 Codex / GPT-5.3-Codex coding, OSWorld, GDPval, CyberGym, and SWE-Lancer scores.",
    },
    {
        "id": "anthropic_sonnet46",
        "name": "Anthropic - Introducing Claude Sonnet 4.6 and system card",
        "url": "https://www.anthropic.com/news/claude-sonnet-4-6",
        "note": "Sonnet 4.6 launch context and Anthropic-reported SWE-bench Verified / Terminal-Bench values.",
    },
    {
        "id": "anthropic_opus47",
        "name": "Anthropic - Introducing Claude Opus 4.7",
        "url": "https://www.anthropic.com/news/claude-opus-4-7",
        "note": "Opus 4.7 launch context and customer benchmark notes.",
    },
    {
        "id": "benchlm_swepro",
        "name": "BenchLM - SWE-bench Pro leaderboard",
        "url": "https://benchlm.ai/benchmarks/swePro",
        "note": "Current SWE-bench Pro leaderboard including Opus 4.7, ChatGPT 5.5, Kimi K2.6, GLM-5.1, ChatGPT 5.3 Codex, MiniMax M2.7, and DeepSeek V4-Pro Max.",
    },
    {
        "id": "kimi_hf",
        "name": "Moonshot AI - Kimi K2.6 Hugging Face model card",
        "url": "https://huggingface.co/moonshotai/Kimi-K2.6",
        "note": "Kimi K2.6 model summary and evaluation table.",
    },
    {
        "id": "deepseek_hf",
        "name": "DeepSeek - DeepSeek-V4-Pro Hugging Face model card",
        "url": "https://huggingface.co/deepseek-ai/DeepSeek-V4-Pro",
        "note": "DeepSeek V4-Pro Max comparison table and evaluation results.",
    },
    {
        "id": "deepseek_transparency",
        "name": "DeepSeek Transparency Center",
        "url": "https://www.deepseek.com/en/transparency/",
        "note": "Confirmed latest DeepSeek V4 release naming and release date.",
    },
    {
        "id": "glm_nvidia",
        "name": "NVIDIA NIM - GLM-5.1 model card",
        "url": "https://build.nvidia.com/z-ai/glm-5.1/modelcard",
        "note": "GLM-5.1 coding, math, and tool-use scores.",
    },
    {
        "id": "minimax_nvidia",
        "name": "NVIDIA NIM - MiniMax M2.7 model card",
        "url": "https://docs.api.nvidia.com/nim/reference/minimaxai-minimax-m2.7",
        "note": "MiniMax M2.7 coding, productivity, tool-use, and ML competition scores.",
    },
    {
        "id": "aa_opus47",
        "name": "Artificial Analysis - Opus 4.7 analysis",
        "url": "https://artificialanalysis.ai/articles/opus-4-7-everything-you-need-to-know/",
        "note": "Opus 4.7 GDPval-AA Elo and Artificial Analysis Intelligence Index context.",
    },
    {
        "id": "hokai_opus47",
        "name": "Hokai - Claude Opus 4.7 benchmark summary",
        "url": "https://hokai.io/hub/models/claude-4.7-opus",
        "note": "Opus 4.7 SWE-bench Verified, SWE-bench Pro, GPQA-Diamond, OSWorld, HLE with tools, and MCP-Atlas summary.",
    },
    {
        "id": "layerlens_opus47_aime",
        "name": "LayerLens - Claude Opus 4.7 on AIME 2026",
        "url": "https://layerlens.ai/blog/claude-opus-4-7-aime2026",
        "note": "Opus 4.7 AIME 2026 independent/aggregated benchmark entry.",
    },
]


def row(model: str, category: str, benchmark: str, score: float, unit: str, source: str, note: str = "") -> dict:
    return {
        "Model": model,
        "Category": category,
        "Benchmark": benchmark,
        "Score": score,
        "Unit": unit,
        "Source": source,
        "Note": note,
    }


ROWS = [
    # ChatGPT 5.5
    row("ChatGPT 5.5", "Coding", "SWE-Bench Pro", 58.6, "% resolved", "openai_gpt55; benchlm_swepro", "Real GitHub issue resolution."),
    row("ChatGPT 5.5", "Coding", "Terminal-Bench 2.0", 82.7, "% accuracy", "openai_gpt55", "Command-line engineering workflows."),
    row("ChatGPT 5.5", "Writing & Knowledge", "GDPval", 84.9, "% wins or ties", "openai_gpt55", "Knowledge-work tasks across occupations."),
    row("ChatGPT 5.5", "Agentic & Other", "OSWorld-Verified", 78.7, "% success", "openai_gpt55", "Computer-use productivity tasks."),
    row("ChatGPT 5.5", "Agentic & Other", "Toolathlon", 55.6, "% pass", "openai_gpt55", "Tool-use benchmark."),
    row("ChatGPT 5.5", "Agentic & Other", "BrowseComp", 84.4, "% pass", "openai_gpt55", "Difficult web browsing and retrieval."),
    row("ChatGPT 5.5", "Math & Reasoning", "FrontierMath Tier 1-3", 51.7, "% pass", "openai_gpt55", "Advanced frontier math."),
    row("ChatGPT 5.5", "Math & Reasoning", "FrontierMath Tier 4", 35.4, "% pass", "openai_gpt55", "Harder frontier math tier."),
    row("ChatGPT 5.5", "Agentic & Other", "CyberGym", 81.8, "% score", "openai_gpt55", "Cybersecurity challenge suite."),
    # ChatGPT 5.3 Codex
    row("ChatGPT 5.3 Codex", "Coding", "SWE-Bench Pro", 56.8, "% resolved", "openai_gpt53_codex; benchlm_swepro", "Real GitHub issue resolution."),
    row("ChatGPT 5.3 Codex", "Coding", "Terminal-Bench 2.0", 77.3, "% accuracy", "openai_gpt53_codex", "Command-line engineering workflows."),
    row("ChatGPT 5.3 Codex", "Agentic & Other", "OSWorld-Verified", 64.7, "% success", "openai_gpt53_codex", "Computer-use productivity tasks."),
    row("ChatGPT 5.3 Codex", "Writing & Knowledge", "GDPval", 70.9, "% wins or ties", "openai_gpt53_codex", "Knowledge-work tasks across occupations."),
    row("ChatGPT 5.3 Codex", "Agentic & Other", "CyberGym", 77.6, "% score", "openai_gpt53_codex", "Cybersecurity challenge suite."),
    row("ChatGPT 5.3 Codex", "Coding", "SWE-Lancer IC Diamond", 81.4, "% score", "openai_gpt53_codex", "Freelance-style software tasks."),
    # Claude Opus 4.7
    row("Claude Opus 4.7", "Coding", "SWE-Bench Pro", 64.3, "% resolved", "benchlm_swepro; hokai_opus47", "Harder SWE-bench variant."),
    row("Claude Opus 4.7", "Coding", "SWE-Bench Verified", 87.6, "% resolved", "hokai_opus47", "Human-validated SWE-bench subset."),
    row("Claude Opus 4.7", "Coding", "Terminal-Bench 2.0", 69.4, "% accuracy", "openai_gpt55", "Reported in OpenAI ChatGPT 5.5 comparison table."),
    row("Claude Opus 4.7", "Math & Reasoning", "GPQA-Diamond", 94.2, "% accuracy", "hokai_opus47", "PhD-level science Q&A."),
    row("Claude Opus 4.7", "Math & Reasoning", "AIME 2026", 90.0, "% accuracy", "layerlens_opus47_aime", "Competition math; independent/aggregated source."),
    row("Claude Opus 4.7", "Math & Reasoning", "FrontierMath Tier 1-3", 43.8, "% pass", "openai_gpt55", "OpenAI comparison table."),
    row("Claude Opus 4.7", "Math & Reasoning", "FrontierMath Tier 4", 22.9, "% pass", "openai_gpt55", "OpenAI comparison table."),
    row("Claude Opus 4.7", "Writing & Knowledge", "GDPval", 80.3, "% wins or ties", "openai_gpt55", "OpenAI comparison table."),
    row("Claude Opus 4.7", "Writing & Knowledge", "GDPval-AA", 1753, "Elo", "aa_opus47", "Artificial Analysis agentic knowledge-work Elo."),
    row("Claude Opus 4.7", "Agentic & Other", "OSWorld-Verified", 78.0, "% success", "openai_gpt55; hokai_opus47", "Computer-use productivity tasks."),
    row("Claude Opus 4.7", "Agentic & Other", "BrowseComp", 79.3, "% pass", "openai_gpt55", "Web browsing and retrieval."),
    row("Claude Opus 4.7", "Agentic & Other", "CyberGym", 73.1, "% score", "openai_gpt55", "Cybersecurity challenge suite."),
    row("Claude Opus 4.7", "Agentic & Other", "MCP-Atlas Public", 77.3, "% pass", "hokai_opus47", "Multi-tool orchestration."),
    row("Claude Opus 4.7", "Agentic & Other", "HLE w/ tools", 54.7, "% pass", "hokai_opus47", "Humanity's Last Exam with tools."),
    # Claude Sonnet 4.6
    row("Claude Sonnet 4.6", "Coding", "SWE-Bench Verified", 79.6, "% resolved", "anthropic_sonnet46", "Anthropic averaged score."),
    row("Claude Sonnet 4.6", "Coding", "Terminal-Bench 2.0", 59.1, "% accuracy", "anthropic_sonnet46", "Anthropic reported default-thinking score."),
    row("Claude Sonnet 4.6", "Writing & Knowledge", "GDPval-AA", 1674, "Elo", "aa_opus47", "Artificial Analysis noted Sonnet 4.6 near the top of GDPval-AA."),
    # MiniMax M2.7
    row("MiniMax M2.7", "Coding", "SWE-Bench Pro", 56.22, "% resolved", "minimax_nvidia; benchlm_swepro", "NVIDIA lists this as SWE-Pro."),
    row("MiniMax M2.7", "Coding", "VIBE-Pro", 55.6, "% score", "minimax_nvidia", "End-to-end project delivery."),
    row("MiniMax M2.7", "Coding", "Terminal-Bench 2.0", 57.0, "% accuracy", "minimax_nvidia", "Complex terminal workflows."),
    row("MiniMax M2.7", "Coding", "NL2Repo", 39.8, "% score", "minimax_nvidia", "Repository-level engineering."),
    row("MiniMax M2.7", "Writing & Knowledge", "GDPval-AA", 1495, "Elo", "minimax_nvidia", "Professional office/domain capability."),
    row("MiniMax M2.7", "Agentic & Other", "Toolathlon", 46.3, "% pass", "minimax_nvidia", "Tool-use benchmark."),
    row("MiniMax M2.7", "Agentic & Other", "MM Claw", 62.7, "% score", "minimax_nvidia", "Complex skills and agent workflows."),
    row("MiniMax M2.7", "Agentic & Other", "MLE Bench Lite", 66.6, "% medal rate", "minimax_nvidia", "ML competition subset."),
    row("MiniMax M2.7", "Coding", "SWE-Bench Multilingual", 76.5, "% resolved", "minimax_nvidia", "Multi-language software tasks."),
    row("MiniMax M2.7", "Coding", "Multi SWE Bench", 52.7, "% resolved", "minimax_nvidia", "Multi-repository SWE tasks."),
    # Kimi K2.6
    row("Kimi K2.6", "Agentic & Other", "HLE w/ tools", 54.0, "% pass", "kimi_hf", "HLE with search/code/web tools."),
    row("Kimi K2.6", "Agentic & Other", "BrowseComp", 83.2, "% pass", "kimi_hf", "Web browsing and retrieval."),
    row("Kimi K2.6", "Agentic & Other", "DeepSearchQA", 92.5, "F1", "kimi_hf", "Search-heavy answer retrieval."),
    row("Kimi K2.6", "Agentic & Other", "Toolathlon", 50.0, "% pass", "kimi_hf", "Tool-use benchmark."),
    row("Kimi K2.6", "Agentic & Other", "OSWorld-Verified", 73.1, "% success", "kimi_hf", "Computer-use productivity tasks."),
    row("Kimi K2.6", "Coding", "Terminal-Bench 2.0", 66.7, "% accuracy", "kimi_hf", "Terminus-2 harness."),
    row("Kimi K2.6", "Coding", "SWE-Bench Pro", 58.6, "% resolved", "kimi_hf; benchlm_swepro", "Harder SWE-bench variant."),
    row("Kimi K2.6", "Coding", "SWE-Bench Multilingual", 76.7, "% resolved", "kimi_hf", "Multi-language software tasks."),
    row("Kimi K2.6", "Coding", "SWE-Bench Verified", 80.2, "% resolved", "kimi_hf", "Averaged over 10 independent runs."),
    row("Kimi K2.6", "Coding", "SciCode", 52.2, "% score", "kimi_hf", "Scientific coding benchmark."),
    row("Kimi K2.6", "Coding", "OJBench Python", 60.6, "% pass", "kimi_hf", "Online-judge Python tasks."),
    row("Kimi K2.6", "Coding", "LiveCodeBench v6", 89.6, "% pass", "kimi_hf", "Live coding benchmark."),
    row("Kimi K2.6", "Math & Reasoning", "HLE", 34.7, "% pass", "kimi_hf", "Humanity's Last Exam without tools."),
    row("Kimi K2.6", "Math & Reasoning", "AIME 2026", 96.4, "% accuracy", "kimi_hf", "Competition math."),
    row("Kimi K2.6", "Math & Reasoning", "HMMT 2026 Feb", 92.7, "% pass", "kimi_hf", "Advanced contest math."),
    row("Kimi K2.6", "Math & Reasoning", "IMO-AnswerBench", 86.0, "% pass", "kimi_hf", "Olympiad-style math answers."),
    row("Kimi K2.6", "Math & Reasoning", "GPQA-Diamond", 90.5, "% accuracy", "kimi_hf", "PhD-level science Q&A."),
    row("Kimi K2.6", "Writing & Knowledge", "GDPval-AA", 1482, "Elo", "deepseek_hf", "DeepSeek comparison table."),
    # DeepSeek
    row("DeepSeek V4-Pro Max", "Math & Reasoning", "MMLU-Pro", 87.5, "% EM", "deepseek_hf", "Difficult multitask knowledge benchmark."),
    row("DeepSeek V4-Pro Max", "Math & Reasoning", "SimpleQA-Verified", 57.9, "% pass", "deepseek_hf", "Factual Q&A."),
    row("DeepSeek V4-Pro Max", "Math & Reasoning", "GPQA-Diamond", 90.1, "% accuracy", "deepseek_hf", "PhD-level science Q&A."),
    row("DeepSeek V4-Pro Max", "Math & Reasoning", "HLE", 37.7, "% pass", "deepseek_hf", "Humanity's Last Exam without tools."),
    row("DeepSeek V4-Pro Max", "Coding", "LiveCodeBench v6", 93.5, "% pass", "deepseek_hf", "Live coding benchmark."),
    row("DeepSeek V4-Pro Max", "Math & Reasoning", "Codeforces", 3206, "rating", "deepseek_hf", "Competitive programming rating."),
    row("DeepSeek V4-Pro Max", "Math & Reasoning", "HMMT 2026 Feb", 95.2, "% pass", "deepseek_hf", "Advanced contest math."),
    row("DeepSeek V4-Pro Max", "Math & Reasoning", "IMO-AnswerBench", 89.8, "% pass", "deepseek_hf", "Olympiad-style math answers."),
    row("DeepSeek V4-Pro Max", "Coding", "Terminal-Bench 2.0", 67.9, "% accuracy", "deepseek_hf", "Complex terminal workflows."),
    row("DeepSeek V4-Pro Max", "Coding", "SWE-Bench Verified", 80.6, "% resolved", "deepseek_hf", "Human-validated SWE-bench subset."),
    row("DeepSeek V4-Pro Max", "Coding", "SWE-Bench Pro", 55.4, "% resolved", "deepseek_hf; benchlm_swepro", "Max reasoning mode."),
    row("DeepSeek V4-Pro Max", "Coding", "SWE-Bench Multilingual", 76.2, "% resolved", "deepseek_hf", "Multi-language software tasks."),
    row("DeepSeek V4-Pro Max", "Agentic & Other", "BrowseComp", 83.4, "% pass", "deepseek_hf", "Web browsing and retrieval."),
    row("DeepSeek V4-Pro Max", "Agentic & Other", "HLE w/ tools", 48.2, "% pass", "deepseek_hf", "HLE with tools."),
    row("DeepSeek V4-Pro Max", "Writing & Knowledge", "GDPval-AA", 1554, "Elo", "deepseek_hf", "Agentic knowledge-work Elo."),
    row("DeepSeek V4-Pro Max", "Agentic & Other", "MCP-Atlas Public", 73.6, "% pass", "deepseek_hf", "Multi-tool orchestration."),
    row("DeepSeek V4-Pro Max", "Agentic & Other", "Toolathlon", 51.8, "% pass", "deepseek_hf", "Tool-use benchmark."),
    row("DeepSeek V4-Pro Max", "Coding", "HumanEval", 76.8, "% pass@1", "deepseek_hf", "Base model table, simple Python tasks."),
    row("DeepSeek V4-Pro Max", "Math & Reasoning", "GSM8K", 92.6, "% EM", "deepseek_hf", "Grade-school math word problems."),
    row("DeepSeek V4-Pro Max", "Math & Reasoning", "MATH", 64.5, "% EM", "deepseek_hf", "Competition-style math problems."),
    # GLM
    row("GLM-5.1", "Coding", "SWE-Bench Pro", 58.4, "% resolved", "glm_nvidia; benchlm_swepro", "Harder SWE-bench variant."),
    row("GLM-5.1", "Coding", "NL2Repo", 42.7, "% score", "glm_nvidia", "Repository-level engineering."),
    row("GLM-5.1", "Coding", "Terminal-Bench 2.0", 63.5, "% accuracy", "glm_nvidia", "Complex terminal workflows."),
    row("GLM-5.1", "Agentic & Other", "CyberGym", 68.7, "% score", "glm_nvidia", "Cybersecurity challenge suite."),
    row("GLM-5.1", "Math & Reasoning", "AIME 2026", 95.3, "% accuracy", "glm_nvidia", "Competition math."),
    row("GLM-5.1", "Math & Reasoning", "HMMT 2025 Nov", 94.0, "% pass", "glm_nvidia", "Advanced contest math."),
    row("GLM-5.1", "Math & Reasoning", "HMMT 2026 Feb", 82.6, "% pass", "glm_nvidia", "Advanced contest math."),
    row("GLM-5.1", "Math & Reasoning", "GPQA-Diamond", 86.2, "% accuracy", "glm_nvidia", "PhD-level science Q&A."),
    row("GLM-5.1", "Math & Reasoning", "IMO-AnswerBench", 83.8, "% pass", "glm_nvidia", "Olympiad-style math answers."),
    row("GLM-5.1", "Agentic & Other", "HLE", 31.0, "% pass", "glm_nvidia", "Humanity's Last Exam without tools."),
    row("GLM-5.1", "Agentic & Other", "HLE w/ tools", 52.3, "% pass", "glm_nvidia", "HLE with tools."),
    row("GLM-5.1", "Agentic & Other", "BrowseComp", 68.0, "% pass", "glm_nvidia", "Web browsing and retrieval."),
    row("GLM-5.1", "Agentic & Other", "BrowseComp w/ context management", 79.3, "% pass", "glm_nvidia", "Web browsing with context management."),
    row("GLM-5.1", "Agentic & Other", "Tool-Decathlon", 40.7, "% pass", "glm_nvidia", "Tool-use benchmark."),
    row("GLM-5.1", "Agentic & Other", "tau3-Bench", 70.6, "% pass", "glm_nvidia", "Complex tool/customer-service style tasks."),
    row("GLM-5.1", "Agentic & Other", "MCP-Atlas Public", 71.8, "% pass", "glm_nvidia", "Multi-tool orchestration."),
    row("GLM-5.1", "Writing & Knowledge", "GDPval-AA", 1535, "Elo", "deepseek_hf", "DeepSeek comparison table."),
]

def guide(
    benchmark: str,
    category: str,
    ladder: str,
    what: str,
    eleven: str,
    top_example: str,
    similar: str,
) -> dict:
    return {
        "Benchmark": benchmark,
        "Category": category,
        "Difficulty Ladder": ladder,
        "What is": what,
        "Explain for 11": eleven,
        "Top-score example": top_example,
        "Similar tests and differences": similar,
    }


BENCHMARK_GUIDE = [
    guide("GSM8K", "Math & Reasoning", "1 - Starter math", "Grade-school arithmetic word problems.", "Story problems: read the sentence, set up the numbers, and calculate the answer.", "multi-step school word problems with clean arithmetic", "Easier than MATH, AIME, HMMT, and IMO-style tests."),
    guide("HumanEval", "Coding", "1 - Starter coding", "Small Python function-writing tasks checked by hidden tests.", "Write one short function, then the computer tries secret examples to see if it works.", "basic Python functions such as list cleanup, string logic, or simple algorithms", "Simpler and older than LiveCodeBench; much smaller than SWE-Bench."),
    guide("SimpleQA-Verified", "Math & Reasoning", "1 - Starter facts", "Short factual questions with verified answers.", "A fact-check quiz: answer clearly and do not make things up.", "straight factual questions where accuracy matters more than style", "Unlike BrowseComp, the answer should not require a long web hunt."),
    guide("MATH", "Math & Reasoning", "2 - School contest math", "Competition-style math problems covering algebra, geometry, counting, and number theory.", "Hard school math puzzles where one missed step ruins the answer.", "contest math that needs several steps, not just arithmetic", "Harder than GSM8K; easier than AIME/HMMT/IMO and FrontierMath."),
    guide("MMLU-Pro", "Math & Reasoning", "2 - Broad knowledge", "A harder version of MMLU with professional and academic multiple-choice questions.", "A big quiz across school and work subjects, but the questions try to trick shallow guessing.", "multi-subject reasoning across science, law, medicine, and humanities", "Broader than GPQA, which focuses on expert science."),
    guide("MMMLU", "Writing & Knowledge", "2 - Multilingual knowledge", "A multilingual MMLU-style benchmark for knowledge across languages.", "The same big school quiz idea, but not only in English.", "knowledge questions in multiple languages", "Related to GMMLU and MILU; these compare language coverage more than coding power."),
    guide("GMMLU English", "Writing & Knowledge", "2 - Multilingual knowledge", "The English slice of a global MMLU-style benchmark.", "The English version of a world knowledge quiz.", "English-language academic and professional knowledge questions", "Use beside low-resource averages to see whether a model is only strong in major languages."),
    guide("GMMLU low-resource average", "Writing & Knowledge", "2 - Multilingual knowledge", "Average performance across lower-resource languages.", "A test of whether the AI still helps in languages with less training data online.", "knowledge questions in languages with fewer public training examples", "More about language fairness than raw reasoning difficulty."),
    guide("MILU average", "Writing & Knowledge", "2 - Multilingual knowledge", "A multilingual language-understanding average.", "Another world-language report card for general knowledge.", "general knowledge and reasoning across languages", "Similar to GMMLU; compare exact datasets before ranking models."),
    guide("LiveCodeBench v6", "Coding", "2 - Current coding", "Fresh programming contest-style tasks designed to reduce memorization.", "A newer coding contest where the model cannot just remember old answers.", "algorithmic coding problems with hidden tests", "More current than HumanEval; not as project-like as SWE-Bench."),
    guide("OJBench Python", "Coding", "2 - Current coding", "Online-judge Python programming tasks.", "Coding contest puzzles where the answer must pass lots of hidden cases.", "Python algorithm puzzles with strict pass/fail tests", "Close to LiveCodeBench, but focused on online-judge style Python."),
    guide("SciCode", "Coding", "3 - Scientific coding", "Scientific programming tasks that need domain formulas and code.", "Write code for science-style problems, not just toy puzzles.", "math/science code that uses formulas, arrays, or simulation logic", "More science-flavored than HumanEval or OJBench."),
    guide("Terminal-Bench 2.0", "Coding", "3 - Real computer work", "Command-line engineering tasks with files, tools, errors, and iteration.", "The AI uses the computer's command box, reads errors, fixes files, and tries again.", "multi-step terminal workflows like installing, testing, debugging, and file edits", "Tests tool operation; SWE-Bench tests real repository bug fixes."),
    guide("SWE-Bench Verified", "Coding", "3 - Real software fixes", "Human-checked GitHub software issues, usually cleaner than the full benchmark.", "A trusted set of broken software tasks where humans checked the task is fair.", "real bug fixes where the model must patch code and pass tests", "Easier/cleaner than SWE-Bench Pro; less broad than Multilingual."),
    guide("SWE-Bench Pro", "Coding", "4 - Hard software fixes", "A harder real-repository software repair benchmark.", "A tough broken app: find the bug, edit the right files, and prove the fix works.", "messy GitHub issues needing code reading, reasoning, and tests", "Harder than Verified; more directly comparable across frontier coding agents."),
    guide("SWE-Bench Multilingual", "Coding", "4 - Multilingual software", "SWE-style tasks across programming languages.", "Like fixing software, but the project might not be Python.", "bugs in JavaScript, Java, Go, Rust, or other language projects", "Different from language benchmarks: this is programming-language diversity."),
    guide("Multi SWE Bench", "Coding", "4 - Multi-repo software", "Software engineering tasks that can span larger or multiple repositories.", "The AI has to understand more than one folder or project area.", "repository-level changes that need broader context", "Close to SWE-Bench, but more about wider codebase navigation."),
    guide("SWE-Lancer IC Diamond", "Coding", "4 - Freelance software work", "Hard software tasks modeled after independent-contractor work.", "A freelance coding job where the final result must satisfy a client-style task.", "larger practical coding jobs with project constraints", "More work-like than HumanEval; related to SWE-Bench but framed as paid tasks."),
    guide("NL2Repo", "Coding", "4 - Build from request", "Turns natural-language requirements into repository-level code changes.", "Someone describes what they want, and the AI has to change the whole project.", "feature work across files from a plain-English request", "More creation-oriented than SWE-Bench, which starts with known issues."),
    guide("VIBE-Pro", "Coding", "4 - Full project delivery", "End-to-end project delivery or app-building benchmark.", "Can the AI build the whole little product, not just solve one puzzle?", "multi-file implementation tasks with polish and working behavior", "Broader than NL2Repo; more product-like than a coding contest."),
    guide("Codeforces", "Math & Reasoning", "4 - Competitive programming", "A competitive-programming rating based on algorithmic problem solving.", "Like a chess rating, but for programming contests.", "advanced algorithms under contest-style constraints", "Comparable to LiveCodeBench only loosely; rating and percent scores are different."),
    guide("GPQA-Diamond", "Math & Reasoning", "4 - Expert science", "Very hard graduate-level science questions.", "Science questions where even smart adults may need the right specialty.", "physics, chemistry, and biology questions that punish guessing", "Narrower and deeper than MMLU-Pro; less math-contest-specific than AIME."),
    guide("AIME 2026", "Math & Reasoning", "4 - Elite school math", "American Invitational Mathematics Examination problems.", "Very hard high-school math contest questions with exact numeric answers.", "number theory, geometry, algebra, and counting puzzles", "Harder than MATH; usually easier than HMMT/IMO and FrontierMath."),
    guide("HMMT 2025 Nov", "Math & Reasoning", "5 - Elite contest math", "Harvard-MIT math tournament problems.", "A top school math team contest with clever, compressed puzzles.", "high-end contest math with several clever insights", "Similar to HMMT 2026 Feb; both sit above AIME-style difficulty."),
    guide("HMMT 2026 Feb", "Math & Reasoning", "5 - Elite contest math", "Harvard-MIT math tournament problems from February 2026.", "A very hard team math contest for students who love puzzles.", "high-end contest math with tricky multi-step solutions", "Harder than AIME; less open-ended than IMO-style answer benchmarks."),
    guide("IMO-AnswerBench", "Math & Reasoning", "5 - Olympiad math", "Olympiad-style math answer evaluation.", "The AI faces math problems closer to world-championship puzzle level.", "proof-like olympiad math where exact reasoning matters", "More proof/olympiad flavored than AIME or HMMT."),
    guide("FrontierMath Tier 1-3", "Math & Reasoning", "6 - Frontier math", "Very hard frontier math problems where current models still miss many.", "Math puzzles so hard that even top AIs do not reliably solve them.", "research-flavored math requiring deep chains of reasoning", "Harder than AIME/HMMT/IMO-style answer sets."),
    guide("FrontierMath Tier 4", "Math & Reasoning", "6 - Frontier math", "The harder FrontierMath tier.", "The boss level of the math ladder in this report.", "the most difficult frontier math items in the source set", "Harder than FrontierMath Tier 1-3."),
    guide("HLE", "Math & Reasoning", "6 - Broad frontier exam", "Humanity's Last Exam: broad, very difficult knowledge and reasoning.", "A giant final exam with questions from many expert fields.", "hard science, humanities, and reasoning questions beyond normal exams", "Broader than GPQA; not limited to math."),
    guide("HLE w/ tools", "Agentic & Other", "6 - Broad frontier with tools", "Humanity's Last Exam with allowed external tools.", "The same giant exam, but the AI can use tools like search or code.", "hard expert questions plus tool choice and research", "Tool access makes it different from plain HLE."),
    guide("ARC-AGI-2 Verified", "Agentic & Other", "6 - Abstract reasoning", "A verified version of ARC-AGI-2, focused on abstract pattern reasoning.", "Colored-grid puzzles where the AI must notice the hidden rule from examples.", "visual pattern puzzles that demand flexible reasoning", "Different from knowledge tests: memorized facts do not help much."),
    guide("GDPval", "Writing & Knowledge", "3 - Professional work", "Professional knowledge-work tasks such as documents, spreadsheets, analysis, and reports.", "Can the AI do office homework well enough that a professional would accept it?", "reports, spreadsheets, legal/finance-style summaries, and analysis", "GDPval is usually percent wins/ties; GDPval-AA uses Elo-style ranking."),
    guide("GDPval-AA", "Writing & Knowledge", "3 - Professional work", "An Artificial Analysis professional-task benchmark using Elo-style scores.", "A league table for office-work ability.", "document, spreadsheet, finance, legal, and business-analysis tasks", "Related to GDPval but the scoring format is different."),
    guide("Finance Agent", "Writing & Knowledge", "4 - Finance work", "Finance-specific agent tasks: analysis, calculations, filings, and tool use.", "A money-and-company homework test for the AI.", "financial analysis workflows, tables, and source-grounded answers", "Narrower than GDPval; more domain-specific."),
    guide("MMMU-Pro", "Writing & Knowledge", "5 - Multimodal expert", "A harder multimodal benchmark using images plus text across expert subjects.", "The AI must read pictures, diagrams, charts, and questions together.", "diagram-heavy science, engineering, and exam questions", "Harder than plain MMMU/MMMLU because vision and expert reasoning combine."),
    guide("MMMU-Pro w/ tools", "Agentic & Other", "5 - Multimodal with tools", "MMMU-Pro with external tools allowed.", "The AI can use tools while solving picture-and-text expert questions.", "multimodal exam problems plus tool-assisted research or calculation", "Different from MMMU-Pro because tool choice becomes part of the test."),
    guide("OSWorld-Verified", "Agentic & Other", "3 - Desktop operation", "Real desktop-computer tasks through a visual interface.", "The AI clicks, types, reads the screen, and completes computer tasks.", "spreadsheet, browser, and app workflows on a computer screen", "More UI-focused than Terminal-Bench, which is command-line focused."),
    guide("BrowseComp", "Agentic & Other", "4 - Web research", "Difficult web browsing and retrieval tasks.", "An internet treasure hunt where the answer is hidden behind several clues.", "web research questions needing persistence and source reading", "More search-heavy than general knowledge tests."),
    guide("BrowseComp w/ context management", "Agentic & Other", "4 - Web research", "BrowseComp with explicit long-context management.", "The AI has to keep track of what it found without forgetting important clues.", "long web investigations with many pieces of evidence", "Harder operationally than plain BrowseComp."),
    guide("DeepSearchQA", "Agentic & Other", "4 - Web research", "Search-heavy question answering scored for answer quality.", "The AI searches and returns the right answer, not just a nice-looking paragraph.", "multi-hop research answers from web evidence", "Similar to BrowseComp but framed as QA quality."),
    guide("Toolathlon", "Agentic & Other", "4 - Tool use", "A benchmark for choosing and using tools correctly.", "The AI has a toolbox and must pick the right tool for each job.", "multi-step tool workflows such as search, calculation, and file actions", "Related to Tool-Decathlon; exact tool set and scoring differ."),
    guide("Tool-Decathlon", "Agentic & Other", "4 - Tool use", "A ten-event tool-use benchmark.", "A decathlon for tools: many different tool events, one overall score.", "diverse tool calls across task types", "Similar to Toolathlon but organized as a ten-part suite."),
    guide("MCP-Atlas Public", "Agentic & Other", "5 - Tool orchestration", "Multi-tool orchestration using Model Context Protocol-style tasks.", "The AI has to coordinate many connected tools without mixing them up.", "multi-step work across connected apps, data, and tools", "More integration-focused than Toolathlon."),
    guide("tau3-Bench", "Agentic & Other", "5 - Customer workflow", "Tool-using conversational tasks with realistic service/workflow constraints.", "The AI handles a customer-style job while following rules and using tools.", "multi-turn service workflows with tool calls and policy constraints", "More conversation/workflow-oriented than MCP-Atlas."),
    guide("MM Claw", "Agentic & Other", "5 - Complex agent skills", "A benchmark for complex skill execution by agents.", "A mixed obstacle course for an AI worker.", "longer tasks that combine planning, tool use, and execution", "Broader but less standardized than single-skill coding benchmarks."),
    guide("MLE Bench Lite", "Agentic & Other", "5 - ML competition", "A lighter machine-learning competition benchmark.", "The AI tries small data-science contests and aims for a medal-like result.", "model training, data handling, and experiment iteration", "Different from coding tests because success depends on data-science judgment."),
    guide("CyberGym", "Agentic & Other", "5 - Cybersecurity", "Cybersecurity challenge tasks.", "A computer-security obstacle course for finding and fixing digital weaknesses.", "security tasks such as vulnerability analysis and challenge solving", "Specialized domain test; not a general writing or math benchmark."),
]

BENCHMARK_EXPLAINERS = [
    (
        f"{item['Benchmark']} ({item['Difficulty Ladder']})",
        f"{item['What is']} Difference note: {item['Similar tests and differences']}",
        item["Explain for 11"],
    )
    for item in BENCHMARK_GUIDE
]

GUIDE_BY_BENCHMARK = {item["Benchmark"]: item for item in BENCHMARK_GUIDE}


def all_benchmark_names() -> list[str]:
    names: list[str] = []
    for item in BENCHMARK_GUIDE:
        names.append(item["Benchmark"])
    for item in ROWS:
        if item["Benchmark"] not in names:
            names.append(item["Benchmark"])
    return names


def guide_for(benchmark: str) -> dict:
    return GUIDE_BY_BENCHMARK.get(
        benchmark,
        {
            "Benchmark": benchmark,
            "Category": "Other",
            "Difficulty Ladder": "Unranked",
            "What is": "A benchmark included in the source table.",
            "Explain for 11": "It is another scoreboard for one kind of AI skill.",
            "Top-score example": "the task type measured by this benchmark",
            "Similar tests and differences": "No extra comparison note is available yet.",
        },
    )


def score_rows_for_benchmark(benchmark: str) -> list[dict]:
    return [item for item in ROWS if item["Benchmark"] == benchmark]


def benchmark_unit(benchmark: str) -> str:
    rows = score_rows_for_benchmark(benchmark)
    return rows[0]["Unit"] if rows else "score"


def top_score_note(benchmark: str) -> str:
    rows = score_rows_for_benchmark(benchmark)
    details = guide_for(benchmark)
    if not rows:
        return "No comparable public score is recorded in this workbook yet. The guide entry is included so you can track it when a model card or leaderboard publishes scores."
    best = max(rows, key=lambda item: item["Score"])
    score = f"{best['Score']:g} {best['Unit']}"
    return f"Top available score here: {best['Model']} at {score}. At that level, the model is solving {details['Top-score example']}."


SIMILARITY_NOTES = [
    (
        "SWE-Bench family",
        "Verified is the cleaner human-checked set. Pro is harder and better for top coding agents. Multilingual checks non-Python codebases. Multi SWE and NL2Repo move closer to broad repository work.",
    ),
    (
        "Coding puzzle family",
        "HumanEval is the small starter set. OJBench and LiveCodeBench are current contest-style coding. SciCode adds science formulas. Codeforces is a rating, not a percent score.",
    ),
    (
        "Math ladder",
        "GSM8K is the easy entry. MATH is contest practice. AIME is elite high-school math. HMMT and IMO-style tasks are harder. FrontierMath is the boss level.",
    ),
    (
        "Knowledge and expert reasoning",
        "MMLU-Pro is broad. GPQA-Diamond is deep expert science. HLE is broad and extreme. SimpleQA checks short factual reliability.",
    ),
    (
        "Professional work",
        "GDPval measures office and professional tasks as wins/ties. GDPval-AA uses Elo. Finance Agent narrows the lens to finance workflows.",
    ),
    (
        "Tool and agent tests",
        "Toolathlon and Tool-Decathlon ask whether tools are chosen correctly. MCP-Atlas asks whether connected tools are orchestrated. tau3-Bench adds conversation and policy constraints.",
    ),
    (
        "Web research",
        "BrowseComp is a hard web treasure hunt. DeepSearchQA is search-heavy question answering. BrowseComp with context management stresses remembering and organizing evidence.",
    ),
    (
        "Multimodal and multilingual",
        "MMMU-Pro adds images and diagrams. MMMU-Pro with tools allows external help. MMMLU, GMMLU, and MILU ask whether knowledge works across languages.",
    ),
]


def source_name(source_id: str) -> str:
    lookup = {s["id"]: s["name"] for s in SOURCES}
    return "; ".join(lookup.get(part.strip(), part.strip()) for part in source_id.split(";"))


def slug(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", text.lower()).strip("_")


def values_for_benchmark(benchmark: str, preferred_unit_contains: str | None = None) -> list[tuple[str, float]]:
    found = {}
    for item in ROWS:
        if item["Benchmark"] == benchmark:
            if preferred_unit_contains and preferred_unit_contains not in item["Unit"]:
                continue
            found[item["Model"]] = item["Score"]
    return [(model, found[model]) for model in MODELS if model in found]


def draw_bar_chart(title: str, values: list[tuple[str, float]], out_path: Path, unit: str = "") -> None:
    width, height = 1200, 720
    margin_left, margin_right = 290, 80
    margin_top, margin_bottom = 105, 80
    plot_w = width - margin_left - margin_right
    img = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(img)
    try:
        font_title = ImageFont.truetype("arial.ttf", 34)
        font = ImageFont.truetype("arial.ttf", 22)
        font_small = ImageFont.truetype("arial.ttf", 18)
    except OSError:
        font_title = ImageFont.load_default()
        font = ImageFont.load_default()
        font_small = ImageFont.load_default()

    draw.text((margin_left, 32), title, fill="#16213e", font=font_title)
    if not values:
        draw.text((margin_left, 160), "No comparable public scores available.", fill="#555555", font=font)
        img.save(out_path)
        return

    max_value = max(v for _, v in values)
    axis_max = 100 if max_value <= 100 else max_value * 1.08
    bar_h = min(50, max(28, int((height - margin_top - margin_bottom) / max(1, len(values)) * 0.62)))
    gap = max(18, int(((height - margin_top - margin_bottom) - len(values) * bar_h) / max(1, len(values))))
    colors = ["#2454A6", "#2A9D8F", "#E76F51", "#6A4C93", "#F4A261", "#3A86FF", "#7B2CBF", "#588157"]

    # Axis grid.
    for tick in range(0, 101, 20):
        x = margin_left + int(plot_w * (tick / 100 if axis_max <= 100 else tick / axis_max))
        draw.line((x, margin_top - 10, x, height - margin_bottom + 12), fill="#e8e8e8", width=1)
        draw.text((x - 12, height - margin_bottom + 25), str(tick), fill="#666666", font=font_small)

    y = margin_top
    for idx, (name, value) in enumerate(values):
        bar_w = int(plot_w * (value / axis_max))
        fill = colors[idx % len(colors)]
        draw.text((20, y + 9), name, fill="#222222", font=font)
        draw.rounded_rectangle((margin_left, y, margin_left + bar_w, y + bar_h), radius=8, fill=fill)
        label = f"{value:g}{unit}"
        draw.text((margin_left + bar_w + 12, y + 9), label, fill="#222222", font=font)
        y += bar_h + gap

    draw.line((margin_left, height - margin_bottom + 10, width - margin_right, height - margin_bottom + 10), fill="#333333", width=2)
    img.save(out_path, quality=95)


def add_table(document: Document, headers: list[str], rows: list[list[object]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, header in enumerate(headers):
        hdr[i].text = header
        for p in hdr[i].paragraphs:
            for run in p.runs:
                run.font.bold = True
    for row_values in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row_values):
            cells[i].text = "" if value is None else str(value)


def format_doc(document: Document) -> None:
    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10)
    styles["Heading 1"].font.name = "Aptos Display"
    styles["Heading 2"].font.name = "Aptos Display"


def create_docx(charts: dict[str, Path]) -> None:
    doc = Document()
    format_doc(doc)
    title = doc.add_heading("Frontier AI Model Benchmark Comparison", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Generated May 24, 2026. Models compared: ChatGPT 5.5, ChatGPT 5.3 Codex, Claude Opus 4.7, Claude Sonnet 4.6, MiniMax M2.7, Kimi K2.6, DeepSeek V4-Pro Max, and GLM-5.1.")

    doc.add_heading("Short Read First", level=1)
    doc.add_paragraph(
        "This revision is benchmark-first: each benchmark keeps its own table and meaning. That matters because a coding repair score, a math contest score, and an office-work Elo score are not interchangeable."
    )
    doc.add_paragraph(
        "Blank cells or missing models mean the source set did not publish a comparable public number. They are not treated as zero."
    )

    doc.add_heading("Model Introductions", level=1)
    for model in MODELS:
        doc.add_heading(model, level=2)
        doc.add_paragraph(MODEL_NOTES[model])

    doc.add_heading("Benchmark Ladder and Plain-English Boxes", level=1)
    doc.add_paragraph(
        "Read the ladder as easiest-to-hardest within a rough family, not as a universal scientific ranking. A tool-use benchmark and a math benchmark can both be hard in different ways."
    )
    guide_rows = []
    for item in BENCHMARK_GUIDE:
        first_col = (
            f"{item['Benchmark']} | {item['Difficulty Ladder']}\n"
            f"What is: {item['What is']}\n"
            f"Difference: {item['Similar tests and differences']}\n"
            f"{top_score_note(item['Benchmark'])}"
        )
        guide_rows.append([first_col, item["Explain for 11"]])
    add_table(doc, ["Benchmark, What Is, and Difference", "Explain It for an 11-Year-Old"], guide_rows)

    doc.add_heading("Similar Tests Cheat Sheet", level=1)
    add_table(doc, ["Family", "How to tell them apart"], SIMILARITY_NOTES)

    doc.add_heading("Score Tables by Benchmark", level=1)
    doc.add_paragraph(
        "Each benchmark below shows only models with sourced public values. The companion XLSX has a Choose Benchmark sheet with a dropdown and a live bar chart."
    )
    for benchmark in all_benchmark_names():
        rows = score_rows_for_benchmark(benchmark)
        if not rows:
            continue
        doc.add_heading(benchmark, level=2)
        doc.add_paragraph(top_score_note(benchmark))
        add_table(
            doc,
            ["Model", "Score", "Unit", "Source"],
            [[item["Model"], item["Score"], item["Unit"], source_name(item["Source"])] for item in rows],
        )

    doc.add_heading("Full Benchmark Table", level=1)
    compact_rows = []
    for item in ROWS:
        compact_rows.append([
            item["Category"],
            item["Benchmark"],
            item["Model"],
            item["Score"],
            item["Unit"],
            source_name(item["Source"]),
        ])
    add_table(doc, ["Category", "Benchmark", "Model", "Score", "Unit", "Source"], compact_rows)

    doc.add_heading("Reading Caveats", level=1)
    for caveat in [
        "Benchmarks are not all run with the same harness, prompting, tool access, or reasoning effort.",
        "Agentic results can move by several points depending on scaffolding and allowed tools.",
        "GDPval and GDPval-AA are related but not the same metric; the report keeps percent win/tie scores and Elo scores separate.",
        "For deployment decisions, rerun the top candidates on your own tasks with the same prompts, tools, budget, and latency constraints.",
    ]:
        doc.add_paragraph(caveat, style=None)

    doc.add_heading("Sources", level=1)
    source_rows = [[s["id"], s["name"], s["url"], s["note"]] for s in SOURCES]
    add_table(doc, ["ID", "Source", "URL", "Used for"], source_rows)

    doc.save(DOCX_PATH)


def autosize(ws) -> None:
    for col in ws.columns:
        max_len = 0
        letter = get_column_letter(col[0].column)
        for cell in col:
            max_len = max(max_len, len(str(cell.value)) if cell.value is not None else 0)
        ws.column_dimensions[letter].width = min(max(max_len + 2, 12), 65)
    for row_cells in ws.iter_rows():
        for cell in row_cells:
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def add_header_style(ws) -> None:
    fill = PatternFill("solid", fgColor="1F4E79")
    for cell in ws[1]:
        cell.fill = fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="top")


def fill_box(ws, start_row: int, label: str, text: str, fill: str = "EAF3F8") -> int:
    ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=9)
    ws.cell(start_row, 1, label).font = Font(bold=True, color="1F4E79")
    ws.cell(start_row, 1).fill = PatternFill("solid", fgColor=fill)
    ws.merge_cells(start_row=start_row + 1, start_column=1, end_row=start_row + 3, end_column=9)
    cell = ws.cell(start_row + 1, 1, text)
    cell.alignment = Alignment(wrap_text=True, vertical="top")
    cell.fill = PatternFill("solid", fgColor="FFFFFF")
    thin = Side(style="thin", color="B7C9D6")
    for row_cells in ws.iter_rows(min_row=start_row, max_row=start_row + 3, min_col=1, max_col=9):
        for box_cell in row_cells:
            box_cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws.row_dimensions[start_row + 1].height = 45
    ws.row_dimensions[start_row + 2].height = 45
    ws.row_dimensions[start_row + 3].height = 45
    return start_row + 5


def create_metadata_sheets(wb: Workbook, benchmarks: list[str]) -> None:
    meta = wb.create_sheet("Benchmark Metadata")
    meta.append([
        "Benchmark",
        "Category",
        "Unit",
        "Difficulty Ladder",
        "What is",
        "Explain for 11",
        "Top-score example",
        "Similar tests and differences",
        "Top score note",
    ])
    for benchmark in benchmarks:
        details = guide_for(benchmark)
        meta.append([
            benchmark,
            details["Category"],
            benchmark_unit(benchmark),
            details["Difficulty Ladder"],
            details["What is"],
            details["Explain for 11"],
            details["Top-score example"],
            details["Similar tests and differences"],
            top_score_note(benchmark),
        ])
    add_header_style(meta)
    autosize(meta)
    meta.freeze_panes = "A2"

    list_ws = wb.create_sheet("Benchmark List")
    for benchmark in benchmarks:
        list_ws.append([benchmark])
    list_ws.sheet_state = "hidden"

    matrix = wb.create_sheet("Benchmark Matrix")
    matrix.append(["Model", *benchmarks])
    for model in MODELS:
        line = [model]
        for benchmark in benchmarks:
            found = [item for item in ROWS if item["Model"] == model and item["Benchmark"] == benchmark]
            line.append(found[0]["Score"] if found else None)
        matrix.append(line)
    add_header_style(matrix)
    autosize(matrix)
    matrix.sheet_state = "hidden"

    source_matrix = wb.create_sheet("Source Matrix")
    source_matrix.append(["Model", *benchmarks])
    for model in MODELS:
        line = [model]
        for benchmark in benchmarks:
            found = [item for item in ROWS if item["Model"] == model and item["Benchmark"] == benchmark]
            line.append(source_name(found[0]["Source"]) if found else "")
        source_matrix.append(line)
    add_header_style(source_matrix)
    autosize(source_matrix)
    source_matrix.sheet_state = "hidden"


def create_choose_sheet(wb: Workbook, benchmarks: list[str]) -> None:
    ws = wb.active
    ws.title = "Choose Benchmark"
    ws.sheet_view.showGridLines = False
    ws["A1"] = "Choose one benchmark"
    ws["A1"].font = Font(size=18, bold=True, color="1F4E79")
    ws["A2"] = "Benchmark"
    ws["B2"] = "SWE-Bench Pro"
    ws["A3"] = "Models are kept down the left side. Pick a benchmark from B2; the table, chart, and explanation boxes update."
    ws["A3"].alignment = Alignment(wrap_text=True)
    ws.merge_cells("A3:H3")

    dv = DataValidation(type="list", formula1=f"='Benchmark List'!$A$1:$A${len(benchmarks)}", allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(ws["B2"])

    headers = ["Model", "Score", "Unit", "Source"]
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(5, col, header)
        cell.fill = PatternFill("solid", fgColor="1F4E79")
        cell.font = Font(color="FFFFFF", bold=True)
    score_last_col = get_column_letter(len(benchmarks) + 1)
    meta_last_row = len(benchmarks) + 1
    for row_idx, model in enumerate(MODELS, start=6):
        ws.cell(row_idx, 1, model)
        ws.cell(
            row_idx,
            2,
            f'=IFERROR(INDEX(\'Benchmark Matrix\'!$B$2:${score_last_col}$9,MATCH(A{row_idx},\'Benchmark Matrix\'!$A$2:$A$9,0),MATCH($B$2,\'Benchmark Matrix\'!$B$1:${score_last_col}$1,0)),"")',
        )
        ws.cell(
            row_idx,
            3,
            f'=IF($B{row_idx}="","",IFERROR(INDEX(\'Benchmark Metadata\'!$C$2:$C${meta_last_row},MATCH($B$2,\'Benchmark Metadata\'!$A$2:$A${meta_last_row},0)),""))',
        )
        ws.cell(
            row_idx,
            4,
            f'=IF($B{row_idx}="","",IFERROR(INDEX(\'Source Matrix\'!$B$2:${score_last_col}$9,MATCH(A{row_idx},\'Source Matrix\'!$A$2:$A$9,0),MATCH($B$2,\'Source Matrix\'!$B$1:${score_last_col}$1,0)),""))',
        )

    chart = BarChart()
    chart.type = "bar"
    chart.style = 10
    chart.title = "Selected benchmark"
    chart.y_axis.title = "Model"
    chart.x_axis.title = "Score"
    data = Reference(ws, min_col=2, min_row=5, max_row=13)
    cats = Reference(ws, min_col=1, min_row=6, max_row=13)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.height = 8
    chart.width = 15
    ws.add_chart(chart, "F5")

    box_formulas = [
        ("What is this benchmark?", f'=IFERROR(INDEX(\'Benchmark Metadata\'!$E$2:$E${meta_last_row},MATCH($B$2,\'Benchmark Metadata\'!$A$2:$A${meta_last_row},0)),"")'),
        ("Explain it for an 11-year-old", f'=IFERROR(INDEX(\'Benchmark Metadata\'!$F$2:$F${meta_last_row},MATCH($B$2,\'Benchmark Metadata\'!$A$2:$A${meta_last_row},0)),"")'),
        ("Top score example", f'=IFERROR(INDEX(\'Benchmark Metadata\'!$I$2:$I${meta_last_row},MATCH($B$2,\'Benchmark Metadata\'!$A$2:$A${meta_last_row},0)),"")'),
        ("Similar tests and differences", f'=IFERROR(INDEX(\'Benchmark Metadata\'!$H$2:$H${meta_last_row},MATCH($B$2,\'Benchmark Metadata\'!$A$2:$A${meta_last_row},0)),"")'),
    ]
    start = 16
    thin = Side(style="thin", color="B7C9D6")
    for label, formula in box_formulas:
        ws.merge_cells(start_row=start, start_column=1, end_row=start, end_column=9)
        ws.cell(start, 1, label).font = Font(bold=True, color="1F4E79")
        ws.cell(start, 1).fill = PatternFill("solid", fgColor="EAF3F8")
        ws.merge_cells(start_row=start + 1, start_column=1, end_row=start + 3, end_column=9)
        ws.cell(start + 1, 1, formula)
        ws.cell(start + 1, 1).alignment = Alignment(wrap_text=True, vertical="top")
        for row_cells in ws.iter_rows(min_row=start, max_row=start + 3, min_col=1, max_col=9):
            for box_cell in row_cells:
                box_cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
        ws.row_dimensions[start + 1].height = 45
        ws.row_dimensions[start + 2].height = 45
        ws.row_dimensions[start + 3].height = 45
        start += 5

    for col, width in {"A": 24, "B": 16, "C": 18, "D": 50, "E": 4, "F": 16, "G": 16, "H": 16, "I": 16}.items():
        ws.column_dimensions[col].width = width
    for row_idx in range(16, start):
        ws.row_dimensions[row_idx].height = 26
    ws.freeze_panes = "A5"


def create_benchmark_charts_sheet(wb: Workbook, benchmarks: list[str]) -> None:
    ws = wb.create_sheet("Benchmark Charts")
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 42
    ws.column_dimensions["F"].width = 18
    start = 1
    for benchmark in benchmarks:
        details = guide_for(benchmark)
        rows = score_rows_for_benchmark(benchmark)
        ws.merge_cells(start_row=start, start_column=1, end_row=start, end_column=9)
        title_cell = ws.cell(start, 1, f"{benchmark} | {details['Difficulty Ladder']}")
        title_cell.font = Font(size=14, bold=True, color="FFFFFF")
        title_cell.fill = PatternFill("solid", fgColor="1F4E79")
        ws.cell(start + 1, 1, "Model")
        ws.cell(start + 1, 2, "Score")
        ws.cell(start + 1, 3, "Unit")
        ws.cell(start + 1, 4, "Source")
        for col in range(1, 5):
            ws.cell(start + 1, col).font = Font(bold=True)
            ws.cell(start + 1, col).fill = PatternFill("solid", fgColor="D9EAF7")

        if rows:
            for row_offset, item in enumerate(rows, start=2):
                target = start + row_offset
                ws.cell(target, 1, item["Model"])
                ws.cell(target, 2, item["Score"])
                ws.cell(target, 3, item["Unit"])
                ws.cell(target, 4, source_name(item["Source"]))
            chart = BarChart()
            chart.type = "bar"
            chart.style = 10
            chart.title = benchmark
            chart.y_axis.title = "Model"
            chart.x_axis.title = benchmark_unit(benchmark)
            data = Reference(ws, min_col=2, min_row=start + 1, max_row=start + 1 + len(rows))
            cats = Reference(ws, min_col=1, min_row=start + 2, max_row=start + 1 + len(rows))
            chart.add_data(data, titles_from_data=True)
            chart.set_categories(cats)
            chart.height = 6
            chart.width = 13
            ws.add_chart(chart, f"F{start + 1}")
            next_row = start + max(len(rows) + 4, 12)
        else:
            ws.cell(start + 2, 1, "No comparable public score in the current source set.")
            next_row = start + 8

        next_row = fill_box(ws, next_row, "What is", details["What is"])
        next_row = fill_box(ws, next_row, "Explain for 11", details["Explain for 11"])
        next_row = fill_box(ws, next_row, "Top-score example", top_score_note(benchmark))
        next_row = fill_box(ws, next_row, "Similar tests and differences", details["Similar tests and differences"])
        start = next_row + 2


def create_xlsx() -> None:
    wb = Workbook()
    benchmarks = all_benchmark_names()
    create_choose_sheet(wb, benchmarks)
    create_metadata_sheets(wb, benchmarks)

    summary = wb.create_sheet("Summary", 1)
    summary.append(["What changed", "Details"])
    summary_rows = [
        ["Benchmark-first layout", "Use Choose Benchmark to pick one benchmark at a time. The models stay down the left side, with a live bar chart."],
        ["OpenAI naming", "The OpenAI rows are labeled ChatGPT 5.5 and ChatGPT 5.3 Codex."],
        ["Plain-English boxes", "Each benchmark has wide explanation boxes: what it is, explain for 11, top-score example, and similar-test differences."],
        ["Missing scores", "Blank cells mean no comparable public value was found in the selected source set, not a zero score."],
    ]
    for line in summary_rows:
        summary.append(line)
    add_header_style(summary)
    autosize(summary)

    long_ws = wb.create_sheet("All Scores")
    long_ws.append(["Category", "Benchmark", "Model", "Score", "Unit", "Source ID", "Source Name", "Note"])
    for item in ROWS:
        long_ws.append([
            item["Category"],
            item["Benchmark"],
            item["Model"],
            item["Score"],
            item["Unit"],
            item["Source"],
            source_name(item["Source"]),
            item["Note"],
        ])
    add_header_style(long_ws)
    autosize(long_ws)
    long_ws.auto_filter.ref = long_ws.dimensions
    long_ws.freeze_panes = "A2"

    create_benchmark_charts_sheet(wb, benchmarks)

    explainer_ws = wb.create_sheet("Benchmark Guide")
    explainer_ws.append([
        "Benchmark",
        "Category",
        "Difficulty Ladder",
        "What is",
        "Explain for 11",
        "Top-score example",
        "Similar tests and differences",
        "Top score note",
    ])
    for item in BENCHMARK_GUIDE:
        explainer_ws.append([
            item["Benchmark"],
            item["Category"],
            item["Difficulty Ladder"],
            item["What is"],
            item["Explain for 11"],
            item["Top-score example"],
            item["Similar tests and differences"],
            top_score_note(item["Benchmark"]),
        ])
    add_header_style(explainer_ws)
    autosize(explainer_ws)
    explainer_ws.column_dimensions["D"].width = 48
    explainer_ws.column_dimensions["E"].width = 48
    explainer_ws.column_dimensions["F"].width = 45
    explainer_ws.column_dimensions["G"].width = 55
    explainer_ws.column_dimensions["H"].width = 55
    explainer_ws.freeze_panes = "A2"

    similar_ws = wb.create_sheet("Similar Tests")
    similar_ws.append(["Family", "How to tell them apart"])
    for line in SIMILARITY_NOTES:
        similar_ws.append(list(line))
    add_header_style(similar_ws)
    autosize(similar_ws)
    similar_ws.column_dimensions["B"].width = 80

    source_ws = wb.create_sheet("Sources")
    source_ws.append(["ID", "Source", "URL", "Used for"])
    for s in SOURCES:
        source_ws.append([s["id"], s["name"], s["url"], s["note"]])
    add_header_style(source_ws)
    autosize(source_ws)

    wb.active = 0
    wb.save(XLSX_PATH)


def main() -> None:
    charts = {
        "coding_swe_bench_pro": CHART_DIR / "coding_swe_bench_pro.png",
        "math_gpqa": CHART_DIR / "math_gpqa_diamond.png",
        "writing_gdpval_aa": CHART_DIR / "writing_gdpval_aa.png",
        "agentic_browsecomp": CHART_DIR / "agentic_browsecomp.png",
    }
    draw_bar_chart("Coding: SWE-Bench Pro", values_for_benchmark("SWE-Bench Pro"), charts["coding_swe_bench_pro"], "%")
    draw_bar_chart("Math & Reasoning: GPQA-Diamond", values_for_benchmark("GPQA-Diamond"), charts["math_gpqa"], "%")
    draw_bar_chart("Writing & Knowledge: GDPval-AA", values_for_benchmark("GDPval-AA"), charts["writing_gdpval_aa"], "")
    draw_bar_chart("Agentic & Other: BrowseComp", values_for_benchmark("BrowseComp"), charts["agentic_browsecomp"], "%")
    create_docx(charts)
    create_xlsx()
    print(DOCX_PATH)
    print(XLSX_PATH)


if __name__ == "__main__":
    main()
